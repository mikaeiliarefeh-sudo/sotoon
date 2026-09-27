"""Fetch OpenRouter model prices for selected providers and write them to Excel.

Prices from the API are USD per token; they are converted to USD per 1M tokens.
"""
import json
import sys
import urllib.request
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PROVIDERS = ["openai", "anthropic", "deepseek", "cohere"]
URL = "https://openrouter.ai/api/v1/models"
OUT = sys.argv[1] if len(sys.argv) > 1 else "openrouter_prices.xlsx"


def per_million(value):
    try:
        v = float(value)
    except (TypeError, ValueError):
        return None
    return round(v * 1_000_000, 4) if v >= 0 else None  # -1 means "variable"


with urllib.request.urlopen(URL, timeout=60) as r:
    models = json.load(r)["data"]

rows = []
for m in models:
    provider = m["id"].split("/")[0]
    if provider not in PROVIDERS:
        continue
    p = m.get("pricing", {})
    rows.append([
        provider,
        m["id"],
        m.get("name"),
        m.get("context_length"),
        per_million(p.get("prompt")),
        per_million(p.get("completion")),
        per_million(p.get("input_cache_read")),
        per_million(p.get("input_cache_write")),
        per_million(p.get("internal_reasoning")),
        float(p["request"]) if p.get("request") not in (None, "") else None,
        float(p["image"]) if p.get("image") not in (None, "") else None,
        float(p["web_search"]) if p.get("web_search") not in (None, "") else None,
    ])
rows.sort(key=lambda r: (PROVIDERS.index(r[0]), r[1]))

headers = [
    "Provider", "Model ID", "Name", "Context (tokens)",
    "Input $/1M", "Output $/1M", "Cache read $/1M", "Cache write $/1M",
    "Reasoning $/1M", "Per request $", "Per image $", "Web search $",
]

wb = Workbook()
sheets = [("All", rows)] + [(pv.capitalize(), [r for r in rows if r[0] == pv]) for pv in PROVIDERS]
first = True
for title, data in sheets:
    ws = wb.active if first else wb.create_sheet()
    first = False
    ws.title = "OpenAI" if title == "Openai" else ("DeepSeek" if title == "Deepseek" else title)
    ws.append(headers)
    for r in data:
        ws.append(r)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    widths = [11, 42, 42, 14, 12, 12, 14, 14, 14, 13, 12, 12]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        row[3].number_format = "#,##0"
        for c in row[4:]:
            c.number_format = "$#,##0.00##"
    ws.append([])
    ws.append([f"Source: {URL} — fetched {date.today().isoformat()}. Empty = not applicable."])

wb.save(OUT)
print(f"{len(rows)} models -> {OUT}")
